# -*- coding: utf-8 -*-
"""
Script to generate the comprehensive, professional Word report for LAB - Exercise 5.11:
Exercise 5.11 - BiCardSource Unit Testing with Stubs & 100% Branch Coverage (code-e3)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_chapter_heading(doc, text):
    h = doc.add_heading(text, level=1)
    run = h.runs[0]
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy Blue
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    return h

def add_section_heading(doc, text):
    h = doc.add_heading(text, level=2)
    run = h.runs[0]
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(29, 78, 216) # Primary Blue
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_body_paragraph(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(31, 41, 55)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(55, 65, 81)
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6") # Light gray
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Remove borders
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(17, 24, 39)

def create_report():
    doc = docx.Document()

    # Set page margins (1 inch = 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = title_p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN - ĐHQG-HCM\nKHOA KỸ THUẬT PHẦN MỀM\n")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(11)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(107, 114, 128)

    title_main = title_p.add_run("BÁO CÁO THỰC HÀNH LAB 06\n")
    title_main.font.name = 'Times New Roman'
    title_main.font.size = Pt(18)
    title_main.font.bold = True
    title_main.font.color.rgb = RGBColor(30, 58, 138)

    sub_title = title_p.add_run("CHUYÊN ĐỀ KIỂM THỬ ĐƠN VỊ VỚI TEST STUBS VÀ ĐỘ BAO PHỦ NHÁNH (BRANCH COVERAGE 100%)\nGIẢI BÀI TẬP EXERCISE 5.11 - LỚP BICARDSOURCE (DESIGNBOOK 3RD EDITION - CODE-E3)")
    sub_title.font.name = 'Times New Roman'
    sub_title.font.size = Pt(12)
    sub_title.font.bold = True
    sub_title.font.color.rgb = RGBColor(79, 70, 229)

    doc.add_paragraph() # Spacer

    # Chapter 1: Tổng quan bài tập và yêu cầu học thuật
    add_chapter_heading(doc, "1. TỔNG QUAN BÀI TẬP EXERCISE 5.11 VÀ YÊU CẦU ĐỀ BÀI")
    add_body_paragraph(doc, "Bài tập 5.11 nằm trong Chương 5 (Unit Testing) thuộc giáo trình 'Introduction to Software Design with Java' (3rd Edition - 2025) của GS. Martin P. Robillard.", "1.1 Bối cảnh học thuật: ")
    add_body_paragraph(doc, "Yêu cầu chính xác của đề bài từ file tài liệu giáo trình (exercises/e-chapter5.md):", "1.2 Yêu cầu đề tài: ")
    add_code_block(doc, 
"""\"Write a test suite for all the public methods of class BiCardSource 
to achieve 100% branch coverage, using stubs to help isolate the behavior to test.\"
(Solution note: \"This solution will reveal a seeded bug. Once you fix the bug, 
the suite should achieve complete branch coverage.\")""")
    
    add_body_paragraph(doc, "Bộ kiểm thử phải thỏa mãn đồng thời 4 mục tiêu cốt lõi:", "1.3 Mục tiêu kỹ thuật cần đạt: ")
    add_body_paragraph(doc, "Sử dụng mẫu thiết kế Test Stub (lớp CardSourceStub) để giả lập nguồn bài, cô lập hoàn toàn lớp BiCardSource mà không cần tạo bộ bài thật (Deck) hay shuffle ngẫu nhiên.", "  1. Kỹ thuật Test Stubs: ")
    add_body_paragraph(doc, "Bao phủ 100% tất cả các nhánh rẽ (decision branches) của toàn bộ các phương thức công khai (draw, isEmpty, constructor, isNull).", "  2. Độ bao phủ 100% Branch Coverage: ")
    add_body_paragraph(doc, "Phát hiện chính xác lỗi logic cài cắm sẵn (seeded bug) trong phương thức draw() của lớp BiCardSource ban đầu.", "  3. Vạch trần Seeded Bug: ")
    add_body_paragraph(doc, "Đồng bộ hóa 100% mã nguồn theo cấu trúc gói mới nhất của Edition 3 (code-e3) gồm e3.chapter4 (CardSource, Card, Rank, Suit) và e3.chapter5 (BiCardSource), loại bỏ hoàn toàn các phụ thuộc cũ của Edition 2.", "  4. Ghép nối chuẩn xác với code-e3: ")

    # Chapter 2: Kiến trúc hệ thống và ghép nối code-e3
    add_chapter_heading(doc, "2. KIẾN TRÚC MÃ NGUỒN VÀ GHÉP NỐI ĐỒNG BỘ VỚI CODE-E3")
    add_body_paragraph(doc, "Trong thư mục tài liệu chính thức 'TAI LIEU/DesignBook-master', mã nguồn được phân định thành các phiên bản khác nhau. Phiên bản Edition 2 cũ sử dụng các package như 'designbook.chapter4.*' với kiểu truy xuất Java Beans (getRank(), getSuit()). Tuy nhiên, trong phiên bản mới nhất Edition 3 (code-e3 - 2025), tác giả Martin P. Robillard đã tái cấu trúc toàn diện:", "2.1 Điểm khác biệt giữa Edition 2 và Edition 3: ")
    add_body_paragraph(doc, "Lớp Card được chuyển sang phong cách thiết kế hiện đại Record-style accessors (phương thức card.rank() và card.suit() thay vì getRank(), getSuit()).", "  - Record-style accessor: ")
    add_body_paragraph(doc, "Lớp Card áp dụng Flyweight design pattern triệt để với phương thức tĩnh Card.get(Rank, Suit) kết hợp bảng lưu trữ bất biến CARDS[suit][rank].", "  - Flyweight Factory: ")
    add_body_paragraph(doc, "Giao diện CardSource thuộc package 'e3.chapter4', bổ sung Null Object pattern tĩnh CardSource.NULL và phương thức mặc định default boolean isNull() { return false; }.", "  - Interface CardSource: ")
    
    add_body_paragraph(doc, "Dưới đây là sơ đồ kiến trúc các gói mã nguồn đã được tổ chức trong thư mục LAB:", "2.2 Phân bổ gói mã nguồn (Package Structure): ")
    add_code_block(doc,
"""LAB/
├── lib/                                      # JUnit 5.10.2 standalone libraries
│   ├── apiguardian-api-1.1.2.jar
│   ├── junit-jupiter-api-5.10.2.jar
│   ├── junit-jupiter-engine-5.10.2.jar
│   ├── junit-platform-commons-1.10.2.jar
│   ├── junit-platform-engine-1.10.2.jar
│   └── opentest4j-1.2.0.jar
├── src/
│   ├── main/java/
│   │   ├── e3/chapter4/                      # Domain Core (code-e3)
│   │   │   ├── CardSource.java               # Interface nguồn bài + Null Object
│   │   │   ├── Card.java                     # Lá bài Flyweight (rank(), suit())
│   │   │   ├── Rank.java                     # Enum ACE..KING
│   │   │   └── Suit.java                     # Enum CLUBS, DIAMONDS, HEARTS, SPADES
│   │   └── e3/chapter5/
│   │       ├── BiCardSource.java             # UUT đã sửa lỗi (Production/Fixed)
│   │       └── BiCardSourceBuggy.java        # Mã gốc chứa seeded bug để đối chứng
│   └── test/java/designbook/solutions5/
│       ├── Exercise11.java                   # Bộ kiểm thử chuẩn (100% Branch Coverage)
│       ├── Exercise11BugDemonstrator.java    # Bộ kiểm nghiệm vạch trần seeded bug
│       ├── TestRunner.java                   # Custom Reflection runner (xUnit isolated)
│       └── CoverageAnalyzerDemo.java         # Trình phân tích ma trận nhánh & CFG
├── pom.xml                                   # Cấu hình Maven (Java 21 LTS)
├── build_and_run.bat / .ps1                  # Script biên dịch và chạy kiểm thử tự động
└── README.md                                 # Hướng dẫn chi tiết & tài liệu học thuật""")

    # Chapter 3: Thiết kế Test Stub
    add_chapter_heading(doc, "3. THIẾT KẾ TEST STUB VÀ CƠ CHẾ CÔ LẬP HÀNH VI")
    add_body_paragraph(doc, "Để kiểm thử một lớp tổng hợp (Composite) như BiCardSource, nếu sử dụng các đối tượng Deck thực tế sẽ phát sinh các vấn đề: phụ thuộc thứ tự lá bài ngẫu nhiên, khó thiết lập trạng thái nguồn rỗng ngay từ đầu, và vi phạm nguyên tắc kiểm thử độc lập (Test Isolation). Do đó, tác giả sử dụng lớp tĩnh nội bộ CardSourceStub:", "3.1 Vai trò của CardSourceStub: ")
    add_code_block(doc,
"""private static class CardSourceStub implements CardSource {

    private Optional<Card> aCard = Optional.empty();

    CardSourceStub() {}

    CardSourceStub(Card pCard) {
        aCard = Optional.of(pCard);
    }

    @Override
    public Card draw() {
        assert !isEmpty() : "Cannot draw from an empty CardSourceStub";
        Card card = aCard.get();
        aCard = Optional.empty();
        return card;
    }

    @Override
    public boolean isEmpty() {
        return !aCard.isPresent();
    }
}""")
    add_body_paragraph(doc, "Lớp Stub sử dụng kiểu dữ liệu java.util.Optional<Card> làm kho lưu trữ nội bộ. Constructor mặc định tạo ra stub rỗng (isEmpty() == true). Constructor có tham số lưu 1 lá bài xác định. Khi phương thức draw() được gọi, lá bài được lấy ra và Optional bị làm rỗng ngay lập tức, chuyển trạng thái của nguồn thành rỗng. Điều này mô phỏng hoàn hảo hành vi rút cạn của một CardSource.", "3.2 Cơ chế hoạt động: ")
    add_body_paragraph(doc, "Trong môi trường kiểm thử thực tế, để đảm bảo tính cô lập giữa các ca kiểm thử (Isolated Test Execution) theo nguyên lý xUnit, bộ kiểm thử sử dụng annotation @BeforeEach để tái tạo mới 3 stub trước mỗi test method: aEmpty, aContainsAce (chứa ACE of CLUBS), và aContainsTwo (chứa TWO of CLUBS).", "3.3 Đảm bảo Test Isolation với @BeforeEach: ")

    # Chapter 4: Phân tích Seeded Bug và Cách Sửa
    add_chapter_heading(doc, "4. VẠCH TRẦN SEEDED BUG VÀ HIỆU CHỈNH CHUẨN XÁC")
    add_body_paragraph(doc, "Theo hợp đồng thiết kế (Design by Contract) được ghi rõ trong JavaDoc của phương thức:", "4.1 Đặc tả thiết kế (Specification): ")
    add_code_block(doc, "/* Draws from source 1 if not empty, and from source 2 if source 1 is empty. */")
    
    add_body_paragraph(doc, "Trong đoạn mã ban đầu của đề bài, tác giả đã cố ý cài cắm một lỗi logic đảo ngược (Inverted Delegation):", "4.2 Mã nguồn lỗi cài cắm (BiCardSourceBuggy): ")
    add_code_block(doc,
"""@Override
public Card draw() {
    if (!aSource1.isEmpty()) {
        return aSource2.draw(); // LỖI: aSource1 còn bài nhưng lại rút từ aSource2!
    }
    else {
        return aSource1.draw(); // LỖI: aSource1 đã rỗng nhưng lại rút từ aSource1!
    }
}""")

    add_body_paragraph(doc, "Khi chạy bộ kiểm thử Exercise11 trên đoạn mã lỗi trên, 3 ca kiểm thử draw() đều lập tức phát hiện sai sót:", "4.3 Biểu hiện thất bại của các ca kiểm thử: ")
    add_body_paragraph(doc, "Khi aSource1 rỗng và aSource2 chứa ACE, điều kiện !aSource1.isEmpty() là false, nhảy vào nhánh else thực hiện aSource1.draw(). Do aSource1 rỗng, phương thức văng ngoại lệ NoSuchElementException / AssertionError! Ca kiểm thử FAILED.", "  - testDraw_FirstEmpty: ")
    add_body_paragraph(doc, "Khi aSource1 chứa ACE và aSource2 rỗng, điều kiện !aSource1.isEmpty() là true, nhảy vào nhánh if thực hiện aSource2.draw(). Do aSource2 rỗng, phương thức cũng văng ngoại lệ! Ca kiểm thử FAILED.", "  - testDraw_SecondEmpty: ")
    add_body_paragraph(doc, "Khi cả hai nguồn đều có bài (aSource1 có ACE, aSource2 có TWO), điều kiện !aSource1.isEmpty() là true, nhảy vào nhánh if thực hiện aSource2.draw(), trả về TWO thay vì ACE. Phép khẳng định assertSame(ACE, drawn) bị thất bại! Ca kiểm thử FAILED.", "  - testDraw_NeitherEmpty: ")

    add_body_paragraph(doc, "Hiệu chỉnh lại phương thức draw() bằng cách gọi đúng nguồn tương ứng với điều kiện rẽ nhánh, đồng thời bổ sung điều kiện tiên quyết assert !isEmpty() theo đúng triết lý lập trình phòng thủ:", "4.4 Mã nguồn sau khi sửa lỗi (BiCardSource chuẩn): ")
    add_code_block(doc,
"""@Override
public Card draw() {
    assert !isEmpty() : "Cannot draw from an empty BiCardSource";
    if (!aSource1.isEmpty()) {
        return aSource1.draw(); // Đúng đặc tả: ưu tiên rút từ nguồn 1
    }
    else {
        return aSource2.draw(); // Nguồn 1 rỗng thì rút từ nguồn 2
    }
}""")

    # Chapter 5: Chứng minh toán học 100% Branch Coverage
    add_chapter_heading(doc, "5. CHỨNG MINH TOÁN HỌC ĐẠT 100% BRANCH COVERAGE")
    add_body_paragraph(doc, "Độ bao phủ nhánh (Branch Coverage / Decision Coverage) đòi hỏi mọi nhánh rẽ (True và False) của tất cả các cấu trúc điều khiển trong phương thức công khai đều phải được thực thi ít nhất một lần.", "5.1 Định nghĩa học thuật: ")
    
    add_section_heading(doc, "5.2 Phân tích phương thức isEmpty()")
    add_body_paragraph(doc, "Câu lệnh thực thi: return aSource1.isEmpty() && aSource2.isEmpty();", "Biểu thức điều kiện: ")
    add_body_paragraph(doc, "Toán tử logic ngắt ngắn mạch '&&' (short-circuit AND) tạo ra bảng chân trị gồm 3 nhánh thực thi độc lập:", "Cơ chế ngắn mạch: ")
    add_body_paragraph(doc, "Được kích hoạt bởi ca kiểm thử testEmpty_True(). Cả 2 vế đều đánh giá true -> Trả về true.", "  1. Nhánh T && T -> TRUE: ")
    add_body_paragraph(doc, "Được kích hoạt bởi ca kiểm thử testEmpty_False_FirstEmpty(). Vế trái true, vế phải false -> Trả về false.", "  2. Nhánh T && F -> FALSE: ")
    add_body_paragraph(doc, "Được kích hoạt bởi hai ca kiểm thử testEmpty_False_SecondEmpty() và testEmpty_False_NeitherEmpty(). Do vế trái aSource1.isEmpty() là false, toán tử && ngắt ngắn mạch ngay lập tức và trả về false mà không cần đánh giá vế phải.", "  3. Nhánh F (Short-circuit) -> FALSE: ")

    add_section_heading(doc, "5.3 Phân tích phương thức draw()")
    add_body_paragraph(doc, "Cấu trúc điều khiển if (!aSource1.isEmpty()) có đúng hai nhánh rẽ chính:", "Cấu trúc rẽ nhánh: ")
    add_body_paragraph(doc, "Được kích hoạt bởi hai ca kiểm thử testDraw_SecondEmpty() và testDraw_NeitherEmpty(). Luồng điều khiển đi vào khối lệnh IF và thực thi return aSource1.draw().", "  1. Nhánh True (!aSource1.isEmpty() == true): ")
    add_body_paragraph(doc, "Được kích hoạt bởi ca kiểm thử testDraw_FirstEmpty(). Luồng điều khiển đi vào khối lệnh ELSE và thực thi return aSource2.draw().", "  2. Nhánh False (!aSource1.isEmpty() == false): ")

    # Table of Branch Coverage
    add_section_heading(doc, "5.4 Bảng ma trận độ bao phủ nhánh toàn diện")
    table = doc.add_table(rows=8, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Test Method", "aSource1", "aSource2", "Nhánh thực thi", "Kết quả mong đợi"]
    hdr_row = table.rows[0]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = title
        set_cell_background(cell, "1E3A8A") # Dark Navy
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("testEmpty_True", "Empty", "Empty", "isEmpty() T && T", "true"),
        ("testEmpty_False_FirstEmpty", "Empty", "Contains ACE", "isEmpty() T && F", "false"),
        ("testEmpty_False_SecondEmpty", "Contains ACE", "Empty", "isEmpty() F (short-circuit)", "false"),
        ("testEmpty_False_NeitherEmpty", "Contains ACE", "Contains TWO", "isEmpty() F (short-circuit)", "false"),
        ("testDraw_FirstEmpty", "Empty", "Contains ACE", "draw() Else branch", "ACE (từ Source 2)"),
        ("testDraw_SecondEmpty", "Contains ACE", "Empty", "draw() If branch", "ACE (từ Source 1)"),
        ("testDraw_NeitherEmpty", "Contains ACE", "Contains TWO", "draw() If branch", "ACE (từ Source 1)")
    ]

    for row_idx, item in enumerate(data):
        row = table.rows[row_idx + 1]
        bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(item):
            cell = row.cells[col_idx]
            cell.text = text
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(31, 41, 55)

    add_body_paragraph(doc, "Bảng trên chứng minh 100% tất cả các nhánh điều kiện và ngắt ngắn mạch trong cả isEmpty() và draw() đều được kích hoạt đầy đủ.", "Kết luận bao phủ: ", italic=True)

    # Chapter 6: Kết quả thực nghiệm và Log
    add_chapter_heading(doc, "6. KẾT QUẢ THỰC NGHIỆM VÀ NHẬT KÝ THỰC THI (TEST LOG)")
    add_body_paragraph(doc, "Dự án đã được biên dịch với tùy chọn '--release 21' và chạy thực tế trên máy tính người dùng bằng TestRunner tự động hóa. Kết quả ghi nhận tại thời điểm thực thi:", "6.1 Kết quả chạy thực tế: ")
    add_code_block(doc,
"""================================================================================
        RUNNING AUTOMATED UNIT TESTS FOR LAB - Exercise 5.11 (EXERCISE 5.11)                  
     BiCardSource Test Suite with Stubs & 100% Branch Coverage (code-e3)        
================================================================================

[TEST SUITE] Exercise11 (BiCardSource Fixed - 100% Branch Coverage)
  [PASS] testDraw_FirstEmpty
  [PASS] testDraw_NeitherEmpty
  [PASS] testDraw_SecondEmpty
  [PASS] testDraw_SequentialTransition
  [PASS] testEmpty_False_FirstEmpty
  [PASS] testEmpty_False_NeitherEmpty
  [PASS] testEmpty_False_SecondEmpty
  [PASS] testEmpty_True
  [PASS] testIsNull

[TEST SUITE] Exercise11BugDemonstrator (Seeded Bug Verification)
  [CONFIRMED SEEDED BUG 1] testDraw_FirstEmpty threw: NoSuchElementException
  [PASS] demonstrateBugInDraw_FirstEmpty
  [CONFIRMED SEEDED BUG 3] testDraw_NeitherEmpty returned: TWO of CLUBS instead of ACE of CLUBS!
  [PASS] demonstrateBugInDraw_NeitherEmpty
  [CONFIRMED SEEDED BUG 2] testDraw_SecondEmpty threw: NoSuchElementException
  [PASS] demonstrateBugInDraw_SecondEmpty
  [PASS] demonstrateEmptyWorks
--------------------------------------------------------------------------------
OVERALL RESULTS: 13 PASSED, 0 FAILED, TOTAL TESTS: 13
================================================================================
TỔNG CỘNG BAO PHỦ NHÁNH: 100% BRANCH COVERAGE (Tất cả nhánh quyết định đều được kích hoạt).
================================================================================
ALL PROCESSES COMPLETED SUCCESSFULLY!""")

    # Chapter 7: Hướng dẫn biên dịch và kiểm thử
    add_chapter_heading(doc, "7. HƯỚNG DẪN BIÊN DỊCH VÀ CHẠY DỰ ÁN")
    add_body_paragraph(doc, "Người dùng và giảng viên có thể kiểm tra dự án dễ dàng thông qua các phương thức sau:", "7.1 Các cách thực thi: ")
    add_body_paragraph(doc, "Nhấp đúp hoặc chạy lệnh 'build_and_run.bat' trong cửa sổ lệnh Windows CMD.", "  - Cách 1 (CMD Script): ")
    add_body_paragraph(doc, "Mở PowerShell tại thư mục LAB và gõ: '.\\build_and_run.ps1'.", "  - Cách 2 (PowerShell): ")
    add_body_paragraph(doc, "Sử dụng lệnh Maven tiêu chuẩn: 'mvn test'.", "  - Cách 3 (Maven): ")
    add_body_paragraph(doc, "Chạy lệnh tiêu chuẩn: 'javac --release 21 -cp \"lib/*\" -d bin src/main/java/e3/*/*.java src/test/java/designbook/solutions5/*.java' sau đó chạy 'java -cp \"bin;lib/*\" designbook.solutions5.TestRunner'.", "  - Cách 4 (Dòng lệnh thủ công): ")

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BAO_CAO_LAB_EXERCISE_11.docx")
    doc.save(output_path)
    print(f"Report successfully generated at: {output_path}")

if __name__ == "__main__":
    create_report()
